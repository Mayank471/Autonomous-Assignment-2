"""Build the CSV lookup tables used by the helpdesk tools.

Outputs (data/structured/), every row carrying the source_doc_id it came from:
  courses.csv    - course code, title, credits, L-T-P, slot, coordinator, email, room (D15 timetable)
  slots.csv      - slot -> day and time, from the published slot grid (D15)
  faculty.csv    - name, designation, department, email (D28-* People directory)
  catalogue.csv  - every approved course: department, code, title, L-T-P, credits, level (D29 ERP)
  teaching.csv   - who taught which course in which semester (D30 faculty web pages)

Usage: python scripts/build_structured.py
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from helpdesk.config import RAW_DIR, SOURCES_CSV, STRUCTURED_DIR  # noqa: E402

COURSE_FIELDS = ["course_code", "course_name", "offered_to", "credits", "ltp", "instructor",
                 "instructor_email", "slot", "room", "semester", "source_doc_id"]
SLOT_FIELDS = ["slot", "day", "time", "cell_text", "semester", "source_doc_id"]
FACULTY_FIELDS = ["name", "designation", "department", "email", "category", "profile_url", "source_doc_id"]
CATALOGUE_FIELDS = ["course_code", "course_title", "erp_department", "dept_code", "course_type", "level",
                    "L", "T", "P", "credits", "source_doc_id"]
TEACHING_FIELDS = ["course_code", "course_title_on_page", "academic_year", "semester", "semester_label",
                   "instructor", "instructor_department", "instructor_email", "context", "source_url",
                   "source_doc_id"]
CODE_RE = re.compile(r"\b([A-Z]{3})\s?(\d{4})\b")


def _s(v):
    return " ".join(str(v).split()) if v is not None else ""


# ---------------------------------------------------------------- D15 courses
def courses_from_d15(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Courses As per Curriculum"]
    header = [_s(c.value) for c in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h}
    semester = "AY 2023-24 Semester I (B.Tech. 1st Year)"
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = _s(r[idx["Course Code"]])
        if not code:
            continue
        credits = r[idx["Credit"]]
        rows.append(dict(
            course_code=code,
            # titles carry the code in brackets, e.g. "Chemistry [CYL1010]"
            course_name=re.sub(r"\s*\[[^\]]*\]\s*", " ", _s(r[idx["Name of Course"]])).strip(),
            offered_to=_s(r[idx["Branch"]]),
            credits=("%g" % credits) if isinstance(credits, (int, float)) else _s(credits),
            ltp=_s(r[idx["LTP"]]),
            instructor=_s(r[idx["Course Coordinator"]]),
            instructor_email=_s(r[idx["eMail"]]),
            slot=_s(r[idx["Slot"]]),
            room=_s(r[idx["Class Room Location"]]),
            semester=semester,
            source_doc_id="D15",
        ))
    return rows


# ------------------------------------------------------------------ D15 slots
def slots_from_d15(path: Path):
    ws = openpyxl.load_workbook(path, data_only=True)["UG 1st Year Slots"]
    semester = "AY 2023-24 Semester I (B.Tech. 1st Year)"
    days = {c.column: _s(c.value) for c in ws[2] if _s(c.value) in {"Mon", "Tue", "Wed", "Thu", "Fri"}}
    times = {}
    for r in range(3, ws.max_row + 1):
        t = _s(ws.cell(r, 1).value)
        if re.match(r"^\d{2}:\d{2}", t):
            times[r] = t
    # merged cells (2-hour labs) span several time rows
    span = {}
    for m in ws.merged_cells.ranges:
        span[(m.min_row, m.min_col)] = m.max_row
    rows = []
    for col, day in days.items():
        for r, t in times.items():
            text = _s(ws.cell(r, col).value)
            if not text:
                continue
            end_row = span.get((r, col), r)
            start = t.split("–")[0].strip()
            end = times.get(end_row, t).split("–")[-1].strip()
            m = re.match(r"^([A-Z]{1,2}\d{0,2})\b", text)
            rows.append(dict(slot=m.group(1) if m else text, day=day, time=f"{start} – {end}",
                             cell_text=text, semester=semester, source_doc_id="D15"))
    return rows


# -------------------------------------------------------------- D28 faculty
def deobfuscate(s: str) -> str:
    return s.replace("[at]", "@").replace("[dot]", ".").replace(" ", "")


def faculty_from_d28(path: Path, doc_id: str):
    text = path.read_text(encoding="utf-8")
    rows = []
    for part in text.split("<!-- SOURCE: ")[1:]:
        soup = BeautifulSoup(part, "html.parser")
        heads = [_s(h.get_text()) for h in soup.find_all(["h1", "h2", "h3"])]
        dept = heads[0] if heads else ""
        category = heads[1] if len(heads) > 1 else ""
        # only faculty listings; staff contact details are not needed by the helpdesk
        if not re.search(r"faculty|^head$", category, re.I):
            continue
        for card in soup.select("div.people-div"):
            name_el = card.select_one(".people-name")
            if not name_el:
                continue
            email = ""
            for info in card.select(".people-contactinfo"):
                icon = info.select_one(".material-icons")
                if icon and _s(icon.get_text()) == "email":
                    icon.extract()
                    email = deobfuscate(_s(info.get_text()))
            desig = card.select_one(".people-designation")
            link = card.find("a", href=True) or card.find_parent("a", href=True)
            rows.append(dict(
                name=_s(name_el.get_text()),
                designation=_s(desig.get_text()) if desig else "",
                department=dept, email=email, category=category,
                profile_url=("https://iitj.ac.in" + link["href"]) if link and link["href"].startswith("/") else "",
                source_doc_id=doc_id,
            ))
    return rows


# ------------------------------------------------------------ D29 catalogue
def catalogue_from_d29(folder: Path):
    """Parse the ERP course cards: title (h4), 'L | T | P | C' and the 'Course Code' label."""
    rows = []
    for page in sorted(folder.glob("*.html")):
        soup = BeautifulSoup(page.read_text(encoding="utf-8"), "html.parser")
        for label in soup.select("span.label"):
            code = _s(label.get_text()).upper().replace(" ", "")
            card = label.find_parent("div", style=re.compile("border"))
            if not card or not re.fullmatch(r"[A-Z]{3}\d{4}", code):
                continue
            heads = card.find_all("h4")
            title = _s(heads[0].get_text()) if heads else ""
            nums = re.findall(r"\d+(?:\.\d+)?", heads[1].get_text()) if len(heads) > 1 else []
            ltpc = nums[-4:] if len(nums) >= 4 else ["", "", "", ""]
            rows.append(dict(course_code=code, course_title=title, erp_department=page.stem,
                             dept_code=code[:2], course_type=code[2], level=int(code[3]) * 100,
                             L=ltpc[0], T=ltpc[1], P=ltpc[2], credits=ltpc[3], source_doc_id="D29"))
    uniq = {}
    for r in rows:  # a course listed under two departments is kept once
        uniq.setdefault(r["course_code"], r)
    return list(uniq.values())


# ---------------------------------------------------------- D30 teaching pages
YEAR_RE = re.compile(r"\b(20\d{2})\s*[-–/]\s*(20\d{2}|\d{2})\b")
SEM_RE = re.compile(r"\b(?:Sem(?:ester)?\.?\s*[-–]?\s*(II|I|2|1)\b|(Odd|Even|Autumn|Fall|Spring|Winter|Summer)\b"
                    r"|(Trimester\s*[-–]?\s*\d))", re.I)
SEASON_YEAR_RE = re.compile(r"\b(Autumn|Fall|Spring|Summer|Winter)\s*[,'-]?\s*(20\d{2})\b", re.I)


def _norm_sem(label: str) -> str:
    """Map only unambiguous labels: I/1/Odd/Autumn/Fall -> I, II/2/Even/Spring -> II, Summer."""
    l = label.lower().replace(" ", "")
    if l in ("i", "1", "odd", "autumn", "fall"):
        return "I"
    if l in ("ii", "2", "even", "spring"):
        return "II"
    if l == "summer":
        return "Summer"
    if l.startswith("trimester"):
        return "Trimester " + l[-1]
    return ""  # e.g. "Winter" is ambiguous; kept only in semester_label


def _page_lines(html: str) -> list:
    """Visible text split into lines: block elements and <br> end a line; inline text stays together."""
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "nav", "header", "footer"]):
        t.decompose()
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for el in soup.find_all(["p", "li", "tr", "div", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"]):
        el.insert_after("\n")
    for td in soup.find_all(["td", "th"]):
        td.insert_after(" | ")
    return [line for line in (_s(x) for x in soup.get_text().split("\n")) if line]


def teaching_from_d30(folder: Path):
    """One row per (course code, academic year, semester, faculty) found on a line of a faculty page.

    A line such as "MAL1010 Mathematics I - Semester I, AY 2026-27, 2022-2023" yields one row per year.
    When the year is on the following lines (card layouts), up to two following lines are used too.
    """
    index = folder / "index.csv"
    if not index.exists():
        return []
    with open(index, encoding="utf-8") as f:
        pages = list(csv.DictReader(f))
    rows, seen = [], set()
    for pg in pages:
        lines = _page_lines((folder / pg["file"]).read_text(encoding="utf-8"))
        for k, line in enumerate(lines):
            if len(line) > 300 or not CODE_RE.search(line):
                continue
            ctx = line
            # card layouts put code, title and "Autumn 2026-27" on consecutive lines: look ahead up to 2 lines
            for nxt in lines[k + 1:k + 3]:
                if YEAR_RE.search(ctx) or SEASON_YEAR_RE.search(ctx) or CODE_RE.search(nxt) or len(nxt) > 120:
                    break
                ctx = f"{ctx} {nxt}"
            years = [f"{a}-{b[-2:]}" for a, b in YEAR_RE.findall(ctx)]
            sm = SEM_RE.search(ctx)
            sem_label = next(g for g in sm.groups() if g) if sm else ""
            sem = _norm_sem(sem_label)
            if not years:  # "Fall 2024" -> AY 2024-25 (Sem I); "Spring 2025" -> AY 2024-25 (Sem II)
                for season, yy in SEASON_YEAR_RE.findall(ctx):
                    ns = _norm_sem(season)
                    if ns in ("I", "II"):
                        start = int(yy) if ns == "I" else int(yy) - 1
                        years.append(f"{start}-{str(start + 1)[-2:]}")
                        sem_label, sem = season, ns
            if not years:
                continue  # an offering without a year cannot be placed in time
            for m in CODE_RE.finditer(line):
                code = m.group(1) + m.group(2)
                after = line[m.end():].lstrip(" :-–)")
                title = re.split(r",|\(|\bSem|\b20\d{2}|\bwith\b|\||;|\bAutumn|\bSpring|\bFall", after, maxsplit=1)[0]
                for year in dict.fromkeys(years):
                    key = (code, year, sem_label.lower(), pg["faculty"])
                    if key in seen:
                        continue
                    seen.add(key)
                    rows.append(dict(course_code=code, course_title_on_page=_s(title)[:80], academic_year=year,
                                     semester=sem, semester_label=sem_label, instructor=pg["faculty"],
                                     instructor_department=pg["department"], instructor_email=pg["faculty_email"],
                                     context=ctx[:300], source_url=pg["url"], source_doc_id="D30"))
    return rows


def write(name, fields, rows):
    STRUCTURED_DIR.mkdir(parents=True, exist_ok=True)
    with open(STRUCTURED_DIR / name, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"{name}: {len(rows)} rows")


def main():
    """Rebuild every table whose raw input is present in data/raw; keep the committed CSV for the others.

    The raw faculty teaching pages (D30) are NOT part of the public repository (they contain student
    names, roll numbers and emails), so data/structured/teaching.csv is kept as committed.
    """
    with open(SOURCES_CSV, encoding="utf-8") as f:
        sources = {r["doc_id"]: r for r in csv.DictReader(f) if r["status"] == "ok"}

    tables = {}  # file name -> (fields, rows)
    skipped = []

    d15 = RAW_DIR / "D15.xlsx"
    if d15.exists():
        tables["courses.csv"] = (COURSE_FIELDS, courses_from_d15(d15))
        tables["slots.csv"] = (SLOT_FIELDS, slots_from_d15(d15))
    else:
        skipped += ["courses.csv", "slots.csv"]

    faculty = []
    for doc_id, src in sources.items():
        path = RAW_DIR.parent.parent / src["local_file"]
        if doc_id.startswith("D28-") and path.exists():
            faculty += faculty_from_d28(path, doc_id)
    if faculty:
        tables["faculty.csv"] = (FACULTY_FIELDS, faculty)
    else:
        skipped.append("faculty.csv")

    erp = RAW_DIR / "D29_erp"
    if "D29" in sources and any(erp.glob("*.html")):
        tables["catalogue.csv"] = (CATALOGUE_FIELDS, catalogue_from_d29(erp))
    else:
        skipped.append("catalogue.csv")

    d30 = RAW_DIR / "D30_teaching"
    if "D30" in sources and any(d30.glob("*.html")):
        tables["teaching.csv"] = (TEACHING_FIELDS, teaching_from_d30(d30))
    else:
        skipped.append("teaching.csv")

    # provenance check: every row must point at a registered source
    for name, (_, rows) in tables.items():
        bad = {r["source_doc_id"] for r in rows} - set(sources)
        if bad:
            raise SystemExit(f"{name}: rows reference unknown source_doc_id {bad}")

    for name, (fields, rows) in tables.items():
        write(name, fields, rows)
    for name in skipped:
        print(f"{name}: raw input not in the repository, keeping the committed file")


if __name__ == "__main__":
    main()
